#!/usr/bin/env python3
"""
fetch_ballot_returns.py
=======================
Downloads (or reads a local copy of) the California SOS VBM ballot return
statistics XLSX and appends the SD4 county figures to ballot_returns.csv.

Usage
-----
  # Download from SOS and append:
  python3 fetch_ballot_returns.py

  # Use a locally-downloaded file instead (e.g. when proxy blocks the CDN):
  python3 fetch_ballot_returns.py --file /path/to/bsr-statistics.xlsx

  # Dry-run — print what would be appended without writing:
  python3 fetch_ballot_returns.py --dry-run

  # Check the undated file AND the dated files (bsr-statistics-MM-DD-YYYY.xlsx)
  # for the last 10 days; append any report dates not already saved:
  python3 fetch_ballot_returns.py --auto

The script is idempotent: if a row for the snapshot date already exists in
ballot_returns.csv it skips rather than duplicating.
"""

import argparse
import csv
import io
import os
import pathlib
import urllib.request
import ssl
import datetime
from typing import Optional

import openpyxl

# ── Paths ──────────────────────────────────────────────────────────────────────
REPO_DIR   = pathlib.Path(__file__).parent
DATA_DIR   = REPO_DIR                          # CSVs live at repo root
OUTPUT_CSV = DATA_DIR / "ballot_returns.csv"
PARTIAL_CSV = DATA_DIR / "partial_county_fractions.csv"

SOS_URL = (
    "https://elections.cdn.sos.ca.gov/"
    "statewide-elections/2026-general/bsr-statistics.xlsx"
)

# ── SD4 counties (all 13) ──────────────────────────────────────────────────────
SD4_COUNTIES = {
    "Alpine", "Amador", "Calaveras", "El Dorado", "Inyo",
    "Madera", "Mariposa", "Merced", "Mono", "Nevada",
    "Placer", "Stanislaus", "Tuolumne",
}

# ── Column indices (0-based) in the data rows ─────────────────────────────────
# Row 4 (index 3) is the header row:
# col 0  COUNTY
# col 1  County Type
# col 2  Total Voters Issued VBM Ballots
# col 3  Drop Box
# col 4  Drop Off Location
# col 5  Vote Center Drop Off
# col 6  Mail
# col 7  FAX
# col 8  Other
# col 9  Total Returned VBM Ballots
# col 10 Total Accepted VBM Ballots
# col 11 Accepted % of Voter-Returned Ballots
COL_COUNTY   = 0
COL_ISSUED   = 2
COL_RETURNED = 9
COL_ACCEPTED = 10


def load_partial_fractions() -> dict:
    fractions = {}
    with open(PARTIAL_CSV, newline="") as f:
        for row in csv.DictReader(f):
            fractions[row["county"]] = float(row["sd4_fraction"])
    return fractions


def dated_urls(d: datetime.date) -> list[str]:
    """SOS also publishes each report under a dated name. The year format has varied:
    bsr-statistics-10-14-2026.xlsx early on, bsr-statistics-10-25-26.xlsx from Oct 25.
    Both are tried. The undated file can lag behind these."""
    return [SOS_URL.replace("bsr-statistics.xlsx", f"bsr-statistics-{d.strftime(fmt)}.xlsx")
            for fmt in ("%m-%d-%y", "%m-%d-%Y")]


def dated_url(d: datetime.date) -> str:
    """Four-digit-year dated name (kept for older callers; prefer dated_urls)."""
    return dated_urls(d)[1]


def is_missing(e: Exception) -> bool:
    """The SOS file server answers 403 (not 404) for a file that doesn't exist."""
    return getattr(e, "code", None) in (403, 404)


def download_xlsx(url: str = SOS_URL) -> bytes:
    """Download an SOS XLSX; raises on failure (including 404)."""
    # Use system CA bundle if available (CCR proxy)
    ca_bundle = "/root/.ccr/ca-bundle.crt"
    if os.path.exists(ca_bundle):
        ctx = ssl.create_default_context(cafile=ca_bundle)
    else:
        ctx = ssl.create_default_context()

    proxy_url = os.environ.get("HTTPS_PROXY") or os.environ.get("https_proxy")
    if proxy_url:
        opener = urllib.request.build_opener(
            urllib.request.ProxyHandler({"https": proxy_url, "http": proxy_url}),
            urllib.request.HTTPSHandler(context=ctx),
        )
    else:
        opener = urllib.request.build_opener(
            urllib.request.HTTPSHandler(context=ctx)
        )

    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (SD4 ballot tracker)"})
    with opener.open(req, timeout=30) as resp:
        return resp.read()


def _num(v) -> Optional[float]:
    """Cell value → number (handles '1,234' strings); None if blank/non-numeric."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _find_layout(ws) -> Optional[dict]:
    """Locate the header row and the county / issued / returned / accepted columns by
    their header text (a header may span two rows). Returns {"row", "county", "issued",
    "returned", "accepted"} (1-based row, 0-based columns) or None."""
    rows = list(ws.iter_rows(min_row=1, max_row=25, values_only=True))
    for r in range(len(rows)):
        prev = rows[r - 1] if r else ()
        width = max(len(rows[r]), len(prev))
        heads = [" ".join(str(x) for x in ((prev[i] if i < len(prev) else None),
                                            (rows[r][i] if i < len(rows[r]) else None)) if x is not None)
                 .lower().replace("\n", " ") for i in range(width)]
        def col(pred):
            hits = [i for i, h in enumerate(heads) if pred(h)]
            return hits
        county = col(lambda h: h.strip().endswith("county") or h.strip() == "county")
        issued = col(lambda h: "issued" in h and "%" not in h)
        returned = col(lambda h: "returned" in h and "%" not in h and "percent" not in h)
        accepted = col(lambda h: "accepted" in h and "%" not in h and "percent" not in h)
        if county and issued and returned:
            ret_total = [i for i in returned if "total" in heads[i]]
            return {"row": r + 1, "county": county[0], "issued": issued[0],
                    "returned": (ret_total or returned)[-1],
                    "accepted": accepted[0] if accepted else None}
    return None


def _describe(wb) -> str:
    """Sheet names and the first rows of each sheet, for the Action log when parsing fails."""
    out = [f"sheets: {wb.sheetnames}"]
    for ws in wb.worksheets[:3]:
        out.append(f"--- {ws.title}")
        for row in ws.iter_rows(min_row=1, max_row=8, max_col=12, values_only=True):
            out.append("    " + " | ".join("" if v is None else str(v)[:25] for v in row))
    return "\n".join(out)


def parse_xlsx(data: bytes, fractions: dict,
               fallback_date: Optional[datetime.date] = None) -> tuple[datetime.date, list[dict]]:
    """
    Parse the SOS XLSX bytes.

    The layout has changed during the 2026 general (sheet name from Oct 24), so the sheet
    and columns are found by name rather than fixed positions:
      sheet   — "VBM Ballot Statistics", else any sheet with "VBM" in its name, else the first
      columns — header row containing County / …Issued… / …Returned… / …Accepted…
      date    — a date/time in the first rows (cell A1 historically), else the file-name date

    Returns
    -------
    snapshot_date : datetime.date
    rows          : list of dicts, one per SD4 county, with keys:
                    county, vbm_issued, vbm_returned, vbm_accepted
                    Values for split counties are already scaled to the SD4 portion.
    """
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    if "VBM Ballot Statistics" in wb.sheetnames:
        candidates = [wb["VBM Ballot Statistics"]]
    else:
        candidates = ([ws for ws in wb.worksheets if "vbm" in ws.title.lower()]
                      + [ws for ws in wb.worksheets if "vbm" not in ws.title.lower()])
    ws = layout = None
    for cand in candidates:
        layout = _find_layout(cand)
        if layout:
            ws = cand
            break
    if not layout:
        raise ValueError("couldn't find the County / Issued / Returned columns\n" + _describe(wb))

    # Report date: first date/time cell above the header row
    snapshot_date = None
    for row in ws.iter_rows(min_row=1, max_row=max(layout["row"] - 1, 1), max_col=12, values_only=True):
        for v in row:
            if isinstance(v, datetime.datetime):
                snapshot_date = v.date()
            elif isinstance(v, datetime.date):
                snapshot_date = v
            if snapshot_date:
                break
        if snapshot_date:
            break
    if snapshot_date is None:
        snapshot_date = fallback_date or datetime.date.today()

    canon = {c.lower(): c for c in SD4_COUNTIES}
    rows_out = []
    for row in ws.iter_rows(min_row=layout["row"] + 1, values_only=True):
        if layout["county"] >= len(row) or row[layout["county"]] is None:
            continue
        county = canon.get(str(row[layout["county"]]).strip().lower())
        if not county:
            continue

        fraction = fractions.get(county, 1.0)
        is_partial = fraction < 1.0

        def get(key):
            i = layout.get(key)
            return _num(row[i]) if i is not None and i < len(row) else None

        def scale(val) -> int:
            return 0 if val is None else round(val * fraction)

        def raw(val) -> Optional[int]:
            """Raw county total for partial counties, else None."""
            return None if (not is_partial or val is None) else round(val)

        iss, ret, acc = get("issued"), get("returned"), get("accepted")
        rows_out.append({
            "county":                county,
            "vbm_issued":            scale(iss),
            "vbm_returned":          scale(ret),
            "vbm_accepted":          scale(acc),
            "county_total_issued":   raw(iss),
            "county_total_returned": raw(ret),
        })

    if not rows_out:
        raise ValueError("found the columns but no SD4 county rows\n" + _describe(wb))
    return snapshot_date, rows_out


def load_existing_keys() -> set:
    """Return set of (county, snapshot_date) already in ballot_returns.csv."""
    if not OUTPUT_CSV.exists():
        return set()
    with open(OUTPUT_CSV, newline="") as f:
        return {(r["county"], r["snapshot_date"]) for r in csv.DictReader(f)}


def append_rows(
    snapshot_date: datetime.date,
    rows: list[dict],
    dry_run: bool = False,
) -> int:
    """Append new rows; return count written (0 if all already present)."""
    existing = load_existing_keys()
    date_str = snapshot_date.isoformat()

    to_write = [
        r for r in rows
        if (r["county"], date_str) not in existing
    ]

    if not to_write:
        print(f"  ✓ All rows for {date_str} already present — nothing to append.")
        return 0

    fieldnames = [
        "county", "snapshot_date",
        "vbm_issued", "vbm_returned", "vbm_accepted",
        "county_total_issued", "county_total_returned",
    ]
    write_header = not OUTPUT_CSV.exists()

    if dry_run:
        print(f"  DRY RUN — would append {len(to_write)} rows for {date_str}:")
        for r in sorted(to_write, key=lambda x: x["county"]):
            ct = (f"  county_total_returned={r['county_total_returned']:>7,}"
                  if r["county_total_returned"] is not None else "")
            print(
                f"    {r['county']:12s}  issued={r['vbm_issued']:>7,}  "
                f"returned={r['vbm_returned']:>6,}  accepted={r['vbm_accepted']:>6,}{ct}"
            )
        return len(to_write)

    with open(OUTPUT_CSV, "a", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        if write_header:
            writer.writeheader()
        for r in sorted(to_write, key=lambda x: x["county"]):
            writer.writerow({
                "county":                r["county"],
                "snapshot_date":         date_str,
                "vbm_issued":            r["vbm_issued"],
                "vbm_returned":          r["vbm_returned"],
                "vbm_accepted":          r["vbm_accepted"],
                "county_total_issued":   r.get("county_total_issued", ""),
                "county_total_returned": r.get("county_total_returned", ""),
            })

    print(f"  ✓ Appended {len(to_write)} rows for {date_str} → {OUTPUT_CSV}")
    return len(to_write)


def fetch_recent(fractions: dict, days_back: int = 10, dry_run: bool = False) -> int:
    """Try the undated file plus each dated file for the last `days_back` days
    and append any report dates not already in ballot_returns.csv.

    Returns the number of new report dates written. Missing files (404) are
    normal — SOS doesn't publish every day — and are skipped quietly.
    """
    today = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=-8))).date()  # ≈ California
    candidates = [(SOS_URL, None)] + [
        (url, today - datetime.timedelta(days=i))
        for i in range(days_back + 1)
        for url in dated_urls(today - datetime.timedelta(days=i))
    ]
    existing_dates = {d for _, d in load_existing_keys()}
    seen, written = set(), 0
    for url, file_date in candidates:
        try:
            data = download_xlsx(url)
        except Exception as e:
            if not is_missing(e):
                print(f"  · {url.rsplit('/', 1)[-1]}: {e}")
            continue
        try:
            snapshot_date, rows = parse_xlsx(data, fractions, fallback_date=file_date)
        except Exception as e:
            print(f"  ✗ {url.rsplit('/', 1)[-1]}: could not parse ({e})")
            continue
        name = url.rsplit("/", 1)[-1]
        if file_date and snapshot_date != file_date:
            print(f"  ! {name}: cell A1 says {snapshot_date}, file name says {file_date}")
        if snapshot_date in seen:
            continue
        seen.add(snapshot_date)
        status = "already have" if snapshot_date.isoformat() in existing_dates else "NEW"
        print(f"  ✓ {name}: report {snapshot_date} · {len(rows)} SD4 counties · {status}")
        if status == "NEW":
            written += 1 if append_rows(snapshot_date, rows, dry_run=dry_run) else 0
    print(f"Done: {written} new report date(s).")
    return written


def main():
    parser = argparse.ArgumentParser(description="Fetch SOS VBM ballot return stats for SD4.")
    parser.add_argument("--file", metavar="PATH", help="Use a local XLSX instead of downloading.")
    parser.add_argument("--auto", action="store_true",
                        help="Check the undated file and the dated files for the last --days days; "
                             "append any report dates not already saved.")
    parser.add_argument("--days", type=int, default=10, help="How many days back --auto looks (default 10).")
    parser.add_argument("--dry-run", action="store_true", help="Print what would be written without writing.")
    args = parser.parse_args()

    fractions = load_partial_fractions()

    if args.auto:
        fetch_recent(fractions, days_back=args.days, dry_run=args.dry_run)
        return

    if args.file:
        print(f"Reading local file: {args.file}")
        with open(args.file, "rb") as f:
            data = f.read()
    else:
        print(f"Downloading from SOS…")
        try:
            data = download_xlsx()
            print("  ✓ Download complete")
        except Exception as e:
            print(f"  ✗ Download failed: {e}")
            print("  Tip: download manually and re-run with --file /path/to/bsr-statistics.xlsx")
            raise SystemExit(1)

    snapshot_date, rows = parse_xlsx(data, fractions)
    print(f"  Snapshot date: {snapshot_date}  |  SD4 counties found: {len(rows)}")

    append_rows(snapshot_date, rows, dry_run=args.dry_run)


if __name__ == "__main__":
    main()
